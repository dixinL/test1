"""Excel 导出：每次运行输出一个 8-sheet 快照。

Sheet1  点位   ：各板块最新交易日收盘点位。
Sheet2~7 MA10..MA360：每窗口一张，全板块结果按 slope_norm 降序。
Sheet8  汇总   ：第二层，每板块一行方向平铺 + 一致性 + UP/DOWN 双榜标注。
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import analyze
import config

# --------- 样式常量 ---------
_HEADER_FONT = Font(bold=True, color='FFFFFF')
_HEADER_FILL = PatternFill('solid', fgColor='4472C4')
_UP_FILL = PatternFill('solid', fgColor='C6EFCE')
_DOWN_FILL = PatternFill('solid', fgColor='FFC7CE')


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'


def _autofit(ws, headers: list[str]) -> None:
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(str(h)) + 4)


def _write_price_sheet(wb: Workbook, snapshot) -> None:
    ws = wb.active
    ws.title = '点位'
    headers = ['行业代码', '行业名称', '交易日', '收盘点位']
    ws.append(headers)
    for r in snapshot:
        ws.append([r['sw_code'], r['sw_name'], r['trade_date'], r['close']])
    _style_header(ws, len(headers))
    _autofit(ws, headers)


def _write_ma_sheets(wb: Workbook, trend_rows) -> None:
    headers = [
        '行业代码', '行业名称', '均线MA', '原始斜率(留档)',
        '归一化斜率', '拟合优度R²', '死区阈值ε', '现价相对均线', '方向',
    ]
    by_window: dict[int, list] = {w: [] for w in config.WINDOWS}
    for r in trend_rows:
        if r['window'] in by_window:
            by_window[r['window']].append(r)

    for w in config.WINDOWS:
        ws = wb.create_sheet(title=f'MA{w}')
        ws.append(headers)
        rows = by_window[w]
        rows.sort(
            key=lambda r: (
                r['slope_norm'] is None,
                -(r['slope_norm'] or 0.0),
            )
        )
        for r in rows:
            ws.append([
                r['sw_code'],
                r['sw_name'],
                _round(r['ma']),
                _round(r['slope'], 8),
                _round(r['slope_norm'], 8),
                _round(r['r_squared'], 4),
                _round(r['epsilon'], 8),
                _round(r['price_vs_ma'], 6),
                r['direction'],
            ])
            _color_direction(ws, ws.max_row, len(headers), r['direction'])
        _style_header(ws, len(headers))
        _autofit(ws, headers)


def _write_summary_sheet(wb: Workbook, summary, up_list, down_list) -> None:
    ws = wb.create_sheet(title='汇总')
    dir_headers = [f'方向_{w}天' for w in config.WINDOWS]
    headers = (
        ['行业代码', '行业名称']
        + dir_headers
        + [
            '综合强度分', '自身历史分位', '上行窗口数',
            '下行窗口数', '一致性', '斜率均值(未加权)', '关注榜',
        ]
    )
    ws.append(headers)

    up_codes = {s['sw_code'] for s in up_list}
    down_codes = {s['sw_code'] for s in down_list}

    def sort_key(s):
        sc = s.get('score', 0.0)
        if s['sw_code'] in up_codes:
            return (0, -sc)
        if s['sw_code'] in down_codes:
            return (1, sc)
        return (2, 0.0)

    for s in sorted(summary, key=sort_key):
        tag = 'UP榜' if s['sw_code'] in up_codes else (
            'DOWN榜' if s['sw_code'] in down_codes else ''
        )
        row = [s['sw_code'], s['sw_name']]
        row += [s[f'dirv_{w}'] for w in config.WINDOWS]
        row += [
            _round(s.get('score'), 8),
            _round(s.get('self_quantile'), 4),
            s['up_count'],
            s['down_count'],
            s['consistency'],
            _round(s['slope_norm_mean'], 8),
            tag,
        ]
        ws.append(row)

        for idx, w in enumerate(config.WINDOWS):
            col_pos = 3 + idx  # 跳过前两列（行业代码、行业名称）
            _color_direction_cell(
                ws.cell(row=ws.max_row, column=col_pos), s[f'dir_{w}']
            )

    _style_header(ws, len(headers))
    _autofit(ws, headers)


def _color_direction(ws, row: int, ncols: int, direction: str) -> None:
    """给最后一列（direction）上色。"""
    _color_direction_cell(ws.cell(row=row, column=ncols), direction)


def _color_direction_cell(cell, direction: str) -> None:
    if direction == 'UP':
        cell.fill = _UP_FILL
    elif direction == 'DOWN':
        cell.fill = _DOWN_FILL


def _round(v, ndigits: int = 4):
    if isinstance(v, (int, float)):
        return round(v, ndigits)
    return v


def export(calc_date: str, snapshot, trend_rows, summary) -> str:
    """写出 Excel 文件，返回文件路径。"""
    os.makedirs(config.EXPORT_DIR, exist_ok=True)

    up_list, down_list = analyze.split_watchlists(summary)

    wb = Workbook()
    _write_price_sheet(wb, snapshot)
    _write_ma_sheets(wb, trend_rows)
    _write_summary_sheet(wb, summary, up_list, down_list)

    path = os.path.join(config.EXPORT_DIR, f'sw_trend_{calc_date}.xlsx')
    wb.save(path)
    return path
